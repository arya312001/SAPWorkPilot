from pathlib import Path

from pypdf import PdfReader


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(str(file_path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if suffix == ".docx":
        try:
            import docx  # python-docx
            doc = docx.Document(str(file_path))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            return file_path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    line = "\t".join(str(c) if c is not None else "" for c in row)
                    if line.strip():
                        lines.append(line)
            return "\n".join(lines)
        except ImportError:
            return ""

    # .txt, .md, .csv and everything else: read as plain text
    return file_path.read_text(encoding="utf-8", errors="ignore")


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 150) -> list[str]:
    """Splits text into overlapping word chunks.

    chunk_size and overlap are word counts, not characters. Overlap keeps a
    sentence that got cut at a chunk boundary readable in the next chunk too,
    so the AI doesn't lose context between chunks.
    """
    words = text.split()
    if not words:
        return []

    chunks = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunk = " ".join(words[start:end])
        if chunk.strip():
            chunks.append(chunk)
        if end == len(words):
            break
        start = end - overlap
    return chunks